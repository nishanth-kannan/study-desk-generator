"""calc blocks: write a quantitative working once; the build shows it as Excel formulae, saves it to an
.xlsx, recalculates it in LibreOffice and checks the result.

    "calc": {"rows": [["Annual demand", 12000], ["Ordering cost", 450], ["Holding cost per unit", 30],
                      ["EOQ", "=SQRT(2*B1*B2/B3)"], ["Orders per year", "=B1/B4"]],
             "answer": ["EOQ", "Orders per year"],     # labels or 1-based rows; default: the last row
             "expect": [600, 20],                      # optional; see below
             "tol": 0.005,                             # relative tolerance, default 0.5%
             "dispute": "The key's 459 is a misprint for 450."}   # only when the official solution disagrees

Row i is cell B<i>; column A holds the label. What each answer is checked against, in order:
  expect (if given)  ->  the correct MCQ option's number  ->  the numbers in e_official.
With none of these the working is shown but nothing is proven (a warning says so).
A mismatch with e_official fails the build unless "dispute" explains it; the dispute is shown in
the explanation as a visible caveat.
"""
import hashlib, json, math, os, re, shutil, subprocess, tempfile
from pathlib import Path

from sdcommon import numbers, plain


def _num(x):
    try: return float(str(x).replace(",", ""))
    except ValueError: return None


def fmt(v):
    if v is None: return ""
    if isinstance(v, str): return v
    if float(v).is_integer() and abs(v) < 1e15: return f"{int(v):,}"
    if abs(v) >= 1000: return f"{v:,.2f}"
    return f"{v:.4g}"


def _close(v, e, tol):
    return v is not None and e is not None and abs(v - e) <= max(tol * abs(e), 1e-9)


def _rows_of(calc, where, ERR):
    rows = calc.get("rows") or []
    if not rows: ERR.append(f"{where}: calc needs 'rows'")
    for i, r in enumerate(rows, 1):
        if not (isinstance(r, list) and len(r) == 2):
            ERR.append(f"{where}: calc row {i} must be [label, value or =formula]")
    return rows


def _answer_rows(calc, rows, where, ERR):
    ans = calc.get("answer", len(rows))
    ans = ans if isinstance(ans, list) else [ans]
    labels = [str(r[0]) for r in rows]
    out = []
    for a in ans:
        if isinstance(a, int) and 1 <= a <= len(rows): out.append(a)
        elif str(a) in labels: out.append(labels.index(str(a)) + 1)
        else: ERR.append(f"{where}: calc answer {a!r} is not a row label or number")
    return out


def recalc(jobs, xlsx_out, cache_file, WARN):
    """jobs: [(sheet, rows)]. Returns {sheet: [values]} and writes the recalculated workbook."""
    import openpyxl
    cache = json.loads(cache_file.read_text()) if cache_file.exists() else {}
    h = lambda rows: hashlib.sha1(json.dumps(rows, sort_keys=True).encode()).hexdigest()
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for sheet, rows in jobs:
        ws = wb.create_sheet(sheet)
        ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 18
        for i, (label, val) in enumerate(rows, 1):
            ws.cell(i, 1, str(label)); ws.cell(i, 2, val)
    need = [s for s, rows in jobs if h(rows) not in cache]
    out = {}
    if need or not xlsx_out.exists():
        if not shutil.which("soffice"):
            WARN.append("LibreOffice (soffice) not found: calc blocks were not recalculated, so nothing was proven")
            wb.save(xlsx_out); return {s: [None] * len(r) for s, r in jobs}
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "calc.xlsx"; wb.save(src)
            prof = Path(td) / "profile"
            subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to",
                            "xlsx:Calc MS Excel 2007 XML", "--outdir", str(Path(td) / "out"), str(src)],
                           check=True, capture_output=True, timeout=180)
            done = Path(td) / "out" / "calc.xlsx"
            vals = openpyxl.load_workbook(done, data_only=True)
            for s, rows in jobs:
                out[s] = [vals[s].cell(i, 2).value for i in range(1, len(rows) + 1)]
                cache[h(rows)] = out[s]
            shutil.copy(done, xlsx_out)
        cache_file.parent.mkdir(parents=True, exist_ok=True)
        cache_file.write_text(json.dumps(cache))
    for s, rows in jobs:
        out.setdefault(s, cache.get(h(rows)))
    return out


def table_html(rows, values, answers, dispute=None):
    body = ""
    for i, (label, val) in enumerate(rows, 1):
        formula = isinstance(val, str) and val.startswith("=")
        shown = f"<code>{val}</code>" if formula else fmt(_num(val) if _num(val) is not None else val)
        res = fmt(values[i - 1]) if values else ""
        star = " class='ans-row'" if i in answers else ""
        body += f"<tr{star}><td>B{i}</td><td>{label}</td><td>{shown}</td><td>{'<b>' + res + '</b>' if i in answers else res}</td></tr>"
    warn = f"<p><span class='warn'>Disputed key: {dispute}</span></p>" if dispute else ""
    return ("<h5>Excel working</h5><figure class='rx'><div class='scroller'><table class='rx-t'><thead><tr><th>Cell</th><th>Item</th>"
            f"<th>Input / formula</th><th>Value</th></tr></thead><tbody>{body}</tbody></table></div></figure>{warn}")


def run(items, out_dir, ERR, WARN, stem):
    """items: [(sheet_key, item, where)]. Checks every calc and appends its table to item['_calc_html']."""
    jobs, meta = [], []
    used = set()
    for key, it, where in items:
        calc = it.get("calc")
        rows = _rows_of(calc, where, ERR)
        if not rows: continue
        sheet = re.sub(r"[\[\]:*?/\\]", "", key)[:31] or "calc"
        base, k = sheet, 2
        while sheet in used: sheet = f"{base[:28]}_{k}"; k += 1
        used.add(sheet)
        jobs.append((sheet, rows)); meta.append((sheet, it, where, rows))
    if not jobs: return None
    xlsx = out_dir / f"{stem} - Excel workings.xlsx"
    values = recalc(jobs, xlsx, out_dir / ".calc_cache.json", WARN)
    proven = 0
    for sheet, it, where, rows in meta:
        calc = it["calc"]
        vals = values.get(sheet) or [None] * len(rows)
        ans = _answer_rows(calc, rows, where, ERR)
        tol = calc.get("tol", 0.005)
        for i, v in enumerate(vals, 1):
            if isinstance(v, str) and v.startswith(("#", "Err")): ERR.append(f"{where}: calc row {i} ({rows[i-1][0]}) gives {v}")
        exp = calc.get("expect")
        exp = exp if isinstance(exp, list) else ([exp] if exp is not None else None)
        source = "expect"
        if exp is None and it.get("o") and isinstance(it.get("a"), int) and not it.get("sub"):
            n = numbers(it["o"][it["a"]])
            if n: exp, source = [_num(n[0])], "the correct option"
        if exp is not None and len(exp) != len(ans):
            ERR.append(f"{where}: calc has {len(ans)} answer rows but {len(exp)} expected values"); continue
        for j, r in enumerate(ans):
            v = vals[r - 1] if r - 1 < len(vals) else None
            v = v if isinstance(v, (int, float)) else _num(v) if v is not None else None
            if v is None and values.get(sheet) and values[sheet][0] is not None: ERR.append(f"{where}: calc row {r} has no numeric value")
            if v is None: continue
            if exp is not None:
                e = _num(exp[j])
                if not (_close(v, e, tol) or _close(v * 100, e, tol)):
                    ERR.append(f"{where}: calc '{rows[r-1][0]}' = {fmt(v)} but {source} says {fmt(e)}")
                else: proven += 1
            elif it.get("e_official"):
                offs = [_num(x) for x in numbers(plain(it["e_official"]))]
                if any(_close(v, o, tol) or _close(v * 100, o, tol) for o in offs): proven += 1
                elif calc.get("dispute"): WARN.append(f"{where}: calc '{rows[r-1][0]}' = {fmt(v)} disagrees with the official solution (disputed)")
                else:
                    ERR.append(f"{where}: calc '{rows[r-1][0]}' = {fmt(v)} is not in the official solution -- fix the working, or add "
                               f"\"dispute\" explaining why the key is wrong")
            else:
                WARN.append(f"{where}: calc '{rows[r-1][0]}' has nothing to check against (add \"expect\")")
        it["_calc_html"] = table_html(rows, vals if values.get(sheet) else None, set(ans), calc.get("dispute"))
    return {"xlsx": xlsx, "sheets": len(jobs), "proven": proven}
